#!/usr/bin/env python3
import argparse
import json
import os
import shutil
import subprocess
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from tempfile import TemporaryDirectory

from delete_rules import (
    compile_delete_rules,
    format_rules_for_summary,
    normalize_cli_delete_rules,
    prune_story_by_rules,
)
from docx import Document


ALLOWED_EXTENSIONS = {".docx", ".doc", ".rtf"}
SOFFICE_FALLBACK_PATHS = [
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/Applications/LibreOffice.app/Contents/MacOS/libreoffice",
    "~/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "~/Applications/LibreOffice.app/Contents/MacOS/libreoffice",
    "/opt/homebrew/bin/soffice",
    "/usr/local/bin/soffice",
]


def detect_soffice_binary() -> str:
    env_path = os.environ.get("DOCSANITIZER_SOFFICE_PATH", "").strip()
    if env_path:
        expanded = Path(env_path).expanduser()
        if expanded.exists() and os.access(expanded, os.X_OK):
            return str(expanded)

    for candidate in ("soffice", "libreoffice"):
        full_path = shutil.which(candidate)
        if full_path:
            return full_path

    for candidate in SOFFICE_FALLBACK_PATHS:
        expanded = Path(candidate).expanduser()
        if expanded.exists() and os.access(expanded, os.X_OK):
            return str(expanded)

    raise RuntimeError(
        "未检测到 LibreOffice，请先安装 libreoffice/soffice。"
        "如果已安装但仍提示未检测到，请设置环境变量 DOCSANITIZER_SOFFICE_PATH"
    )


def sanitize_story(story) -> None:
    elements = []
    for paragraph in story.paragraphs:
        elements.append(paragraph._element)
    for table in story.tables:
        elements.append(table._element)
    for element in elements:
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)
    story.add_paragraph("")


def strip_header_footer_docx(source_path: Path, target_path: Path, compiled_rules) -> None:
    document = Document(str(source_path))
    for section in document.sections:
        stories = [
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ]
        for story in stories:
            prune_story_by_rules(story, compiled_rules, fallback_clear=sanitize_story)
    target_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(target_path))


def libreoffice_convert(
    input_file: Path,
    out_dir: Path,
    convert_to: str,
    timeout_sec: int,
    soffice_bin: str,
    user_profile_dir: Path,
) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    user_profile_dir.mkdir(parents=True, exist_ok=True)
    command = [
        soffice_bin,
        "--headless",
        f"-env:UserInstallation={user_profile_dir.resolve().as_uri()}",
        "--convert-to",
        convert_to,
        "--outdir",
        str(out_dir),
        str(input_file),
    ]
    try:
        subprocess.run(
            command,
            check=True,
            timeout=timeout_sec,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
        )
    except subprocess.CalledProcessError as exc:
        stdout = (exc.stdout or "").strip()
        stderr = (exc.stderr or "").strip()
        raise RuntimeError(
            f"LibreOffice 转换失败({input_file.name} -> {convert_to})，"
            f"exit={exc.returncode}，stdout={stdout[:500]}，stderr={stderr[:500]}"
        ) from exc
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError(
            f"LibreOffice 转换超时({input_file.name} -> {convert_to})，timeout={timeout_sec}s"
        ) from exc

    target_ext = convert_to.split(":")[0].lower()
    expected = out_dir / f"{input_file.stem}.{target_ext}"
    if expected.exists():
        return expected

    candidates = sorted(out_dir.glob(f"{input_file.stem}.*"), key=lambda x: x.stat().st_mtime)
    if not candidates:
        generated = ", ".join(path.name for path in out_dir.glob("*"))
        raise RuntimeError(
            f"文件转换失败：{input_file.name} -> {convert_to}，输出目录为空或无目标文件。"
            f"目录内容: {generated or '空'}"
        )
    return candidates[-1]


def sanitize_document(
    source_path: Path,
    target_path: Path,
    retries: int,
    timeout_sec: int,
    soffice_bin: str,
    delete_rules: list[dict] | None = None,
) -> None:
    compiled_rules = compile_delete_rules(delete_rules or [])
    ext = source_path.suffix.lower()
    if ext == ".docx":
        strip_header_footer_docx(source_path, target_path, compiled_rules)
        return
    if ext not in {".doc", ".rtf"}:
        raise RuntimeError(f"不支持的文件类型: {ext}")

    last_error = None
    for _ in range(retries + 1):
        try:
            with TemporaryDirectory() as working_dir:
                working = Path(working_dir)
                profile_dir = working / "lo_profile"
                source_docx = libreoffice_convert(
                    source_path,
                    working,
                    "docx",
                    timeout_sec,
                    soffice_bin,
                    profile_dir,
                )
                cleaned_docx = working / "cleaned.docx"
                strip_header_footer_docx(source_docx, cleaned_docx, compiled_rules)

                convert_to = 'doc:"MS Word 97"' if ext == ".doc" else "rtf"
                converted = libreoffice_convert(
                    cleaned_docx,
                    working,
                    convert_to,
                    timeout_sec,
                    soffice_bin,
                    profile_dir,
                )
                target_path.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(converted, target_path)
            return
        except Exception as exc:  # noqa: BLE001
            last_error = exc
            time.sleep(0.5)
    raise RuntimeError(f"{source_path.name} 处理失败: {last_error}") from last_error


def collect_input_files(input_dir: Path, recursive: bool) -> list[Path]:
    if not input_dir.exists():
        raise RuntimeError(f"输入目录不存在: {input_dir}")
    if not input_dir.is_dir():
        raise RuntimeError(f"输入路径不是目录: {input_dir}")

    matched: list[Path] = []
    iterator = input_dir.rglob("*") if recursive else input_dir.glob("*")
    for path in iterator:
        if path.is_file() and path.suffix.lower() in ALLOWED_EXTENSIONS:
            matched.append(path)
    return sorted(matched)


def build_target_path(input_root: Path, source_file: Path, output_root: Path) -> Path:
    relative = source_file.relative_to(input_root)
    return output_root / relative.parent / f"{source_file.stem}_sanitized{source_file.suffix.lower()}"


def process_one(
    input_root: Path,
    output_root: Path,
    source_file: Path,
    overwrite: bool,
    retries: int,
    timeout_sec: int,
    soffice_bin: str,
    delete_rules: list[dict],
) -> dict:
    started_at = time.time()
    target_file = build_target_path(input_root, source_file, output_root)
    if target_file.exists() and not overwrite:
        return {
            "file": str(source_file),
            "output": str(target_file),
            "status": "skipped",
            "duration_sec": round(time.time() - started_at, 3),
            "message": "已存在同名输出文件，使用 --overwrite 可覆盖",
        }

    try:
        sanitize_document(source_file, target_file, retries, timeout_sec, soffice_bin, delete_rules=delete_rules)
        return {
            "file": str(source_file),
            "output": str(target_file),
            "status": "success",
            "duration_sec": round(time.time() - started_at, 3),
            "message": "",
        }
    except Exception as exc:  # noqa: BLE001
        return {
            "file": str(source_file),
            "output": str(target_file),
            "status": "failed",
            "duration_sec": round(time.time() - started_at, 3),
            "message": str(exc),
        }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="本地批量删除 .docx/.doc/.rtf 页眉页脚",
        allow_abbrev=False,
    )
    parser.add_argument("--input-dir", required=True, help="输入目录")
    parser.add_argument("--output-dir", required=True, help="输出目录")
    parser.add_argument("--workers", type=int, default=min(4, (os.cpu_count() or 2)), help="并发数")
    parser.add_argument("--recursive", action="store_true", help="是否递归扫描子目录")
    parser.add_argument("--overwrite", action="store_true", help="覆盖已存在输出文件")
    parser.add_argument("--retries", type=int, default=1, help="单文件失败后的重试次数")
    parser.add_argument("--timeout-sec", type=int, default=600, help="单次 LibreOffice 转换超时秒数")
    parser.add_argument(
        "--delete-rule",
        action="append",
        default=[],
        help="删除规则，格式为 mode:value，例如 contains:方案编号、regex:方案编号[:：].*",
    )
    parser.add_argument("--delete-rule-json", default="", help="删除规则 JSON 数组字符串")
    parser.add_argument("--delete-rule-file", default="", help="删除规则 JSON 文件路径")
    parser.add_argument("--delete-pattern", action="append", default=[], help="兼容参数：等同 regex 规则")
    parser.add_argument(
        "--manifest-name",
        default="sanitize_manifest.json",
        help="输出结果清单文件名（写入 output-dir）",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    input_dir = Path(args.input_dir).resolve()
    output_dir = Path(args.output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    files = collect_input_files(input_dir, args.recursive)
    total = len(files)
    if total == 0:
        print("未找到可处理文件（仅支持 .docx/.doc/.rtf）。")
        return 1

    soffice_bin = detect_soffice_binary()
    workers = max(1, args.workers)
    delete_rules = normalize_cli_delete_rules(
        delete_rule_json=args.delete_rule_json,
        delete_rule_file=args.delete_rule_file,
        delete_patterns=[*args.delete_rule, *args.delete_pattern],
    )
    compile_delete_rules(delete_rules)
    started_at = time.time()
    print(f"找到 {total} 个文件，开始处理。并发数={workers}")
    if delete_rules:
        print("启用删除规则:")
        for index, summary in enumerate(format_rules_for_summary(delete_rules), start=1):
            print(f"  {index}. {summary}")
    else:
        print("未配置删除规则：将清空页眉页脚全部内容。")

    results: list[dict] = []
    success = failed = skipped = 0
    done = 0

    with ThreadPoolExecutor(max_workers=workers) as executor:
        futures = [
            executor.submit(
                process_one,
                input_dir,
                output_dir,
                file_path,
                args.overwrite,
                max(0, args.retries),
                max(60, args.timeout_sec),
                soffice_bin,
                delete_rules,
            )
            for file_path in files
        ]

        for future in as_completed(futures):
            done += 1
            result = future.result()

            status = result["status"]
            if status == "success":
                success += 1
            elif status == "skipped":
                skipped += 1
            else:
                failed += 1
            results.append(result)

            print(
                f"[{done}/{total}] {status.upper()} | "
                f"{Path(result['file']).name} -> {Path(result['output']).name if result['output'] else '-'}"
            )
            if result["message"]:
                print(f"           {result['message']}")

    elapsed = round(time.time() - started_at, 3)
    manifest = {
        "summary": {
            "total": total,
            "success": success,
            "failed": failed,
            "skipped": skipped,
            "elapsed_sec": elapsed,
            "workers": workers,
            "delete_rules": delete_rules,
            "input_dir": str(input_dir),
            "output_dir": str(output_dir),
        },
        "results": sorted(results, key=lambda x: x["file"]),
    }

    manifest_path = output_dir / args.manifest_name
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    print("\n处理完成")
    print(f"- 成功: {success}")
    print(f"- 失败: {failed}")
    print(f"- 跳过: {skipped}")
    print(f"- 耗时: {elapsed}s")
    print(f"- 清单: {manifest_path}")
    return 0 if failed == 0 else 2


if __name__ == "__main__":
    raise SystemExit(main())
