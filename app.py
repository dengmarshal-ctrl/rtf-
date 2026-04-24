import json
import re
import shutil
import subprocess
import threading
import time
import uuid
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory

from delete_rules import compile_delete_rules, normalize_delete_rules
from docx import Document
from flask import Flask, abort, jsonify, render_template, request, send_file
from werkzeug.utils import secure_filename


BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
UPLOAD_DIR = DATA_DIR / "uploads"
OUTPUT_DIR = DATA_DIR / "outputs"
TMP_DIR = DATA_DIR / "tmp"

ALLOWED_EXTENSIONS = {".docx", ".doc", ".rtf"}
MAX_FILE_SIZE = 2 * 1024 * 1024 * 1024
MAX_CHUNK_SIZE = 16 * 1024 * 1024

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 32 * 1024 * 1024

_upload_lock = threading.RLock()
_job_lock = threading.RLock()
_uploads = {}
_jobs = {}


def ensure_dirs() -> None:
    for folder in (DATA_DIR, UPLOAD_DIR, OUTPUT_DIR, TMP_DIR):
        folder.mkdir(parents=True, exist_ok=True)


def allowed_file(filename: str) -> bool:
    return Path(filename).suffix.lower() in ALLOWED_EXTENSIONS


def mask_filename(filename: str) -> str:
    path = Path(filename)
    stem = path.stem
    ext = path.suffix
    if not stem:
        return f"***{ext}"
    if len(stem) <= 2:
        return f"{stem[0]}*{ext}" if len(stem) == 2 else f"*{ext}"
    return f"{stem[0]}{'*' * (len(stem) - 2)}{stem[-1]}{ext}"


def sanitize_story(story) -> None:
    elements = []
    for paragraph in story.paragraphs:
        elements.append(paragraph._element)
    for table in story.tables:
        elements.append(table._element)
    for element in elements:
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)
    story.add_paragraph("")


def _remove_paragraph(paragraph) -> None:
    parent = paragraph._element.getparent()
    if parent is not None:
        parent.remove(paragraph._element)


def _remove_table(table) -> None:
    parent = table._element.getparent()
    if parent is not None:
        parent.remove(table._element)


def _match_text(text: str, compiled_rules: list) -> bool:
    for rule in compiled_rules:
        rule_type = rule["type"]
        value = rule["value"]
        text_value = text or ""
        if rule_type == "regex":
            if value.search(text_value):
                return True
        elif rule_type == "contains":
            if value in text_value:
                return True
        elif rule_type == "prefix":
            if text_value.startswith(value):
                return True
        elif rule_type == "suffix":
            if text_value.endswith(value):
                return True
    return False


def sanitize_story_by_rules(story, compiled_rules: list) -> None:
    if not compiled_rules:
        sanitize_story(story)
        return

    for paragraph in list(story.paragraphs):
        if _match_text(paragraph.text, compiled_rules):
            _remove_paragraph(paragraph)

    for table in list(story.tables):
        for row in table.rows:
            for cell in row.cells:
                for paragraph in list(cell.paragraphs):
                    if _match_text(paragraph.text, compiled_rules):
                        _remove_paragraph(paragraph)

        remaining = "".join(
            paragraph.text.strip()
            for row in table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
        )
        if not remaining:
            _remove_table(table)

    if not story.paragraphs and not story.tables:
        story.add_paragraph("")


def strip_header_footer_docx(source_path: Path, target_path: Path, compiled_rules: list) -> None:
    document = Document(str(source_path))
    for section in document.sections:
        stories = [
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ]
        for story in stories:
            sanitize_story_by_rules(story, compiled_rules)
    document.save(str(target_path))


def detect_soffice_binary() -> str:
    for candidate in ("soffice", "libreoffice"):
        full_path = shutil.which(candidate)
        if full_path:
            return full_path
    raise RuntimeError("未检测到 LibreOffice，请安装 libreoffice/soffice 后重试。")


def libreoffice_convert(input_file: Path, out_dir: Path, convert_to: str) -> Path:
    binary = detect_soffice_binary()
    out_dir.mkdir(parents=True, exist_ok=True)
    command = [
        binary,
        "--headless",
        "--convert-to",
        convert_to,
        "--outdir",
        str(out_dir),
        str(input_file),
    ]
    subprocess.run(
        command,
        check=True,
        timeout=600,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )

    target_ext = convert_to.split(":")[0].lower()
    expected = out_dir / f"{input_file.stem}.{target_ext}"
    if expected.exists():
        return expected

    candidates = sorted(out_dir.glob(f"{input_file.stem}.*"), key=lambda x: x.stat().st_mtime)
    if not candidates:
        raise RuntimeError(f"文件转换失败：{input_file.name} -> {convert_to}")
    return candidates[-1]


def sanitize_document(source_path: Path, original_name: str, destination_dir: Path, delete_rules: list[dict]) -> Path:
    destination_dir.mkdir(parents=True, exist_ok=True)
    ext = Path(original_name).suffix.lower()
    safe_stem = Path(secure_filename(original_name)).stem or "document"
    target_name = f"{safe_stem}_sanitized{ext}"
    target_path = destination_dir / target_name
    compiled_rules = compile_delete_rules(delete_rules)

    if ext == ".docx":
        strip_header_footer_docx(source_path, target_path, compiled_rules)
        return target_path

    if ext not in {".doc", ".rtf"}:
        raise RuntimeError(f"不支持的文件类型: {ext}")

    with TemporaryDirectory(dir=str(TMP_DIR)) as working_dir:
        working = Path(working_dir)
        source_docx = libreoffice_convert(source_path, working, "docx")
        cleaned_docx = working / "cleaned.docx"
        strip_header_footer_docx(source_docx, cleaned_docx, compiled_rules)

        convert_to = 'doc:"MS Word 97"' if ext == ".doc" else "rtf"
        converted = libreoffice_convert(cleaned_docx, working, convert_to)
        shutil.move(str(converted), target_path)

    return target_path


def get_upload_chunk_path(upload_info: dict, chunk_index: int) -> Path:
    return Path(upload_info["upload_path"]) / "chunks" / f"{chunk_index:08d}.part"


def build_job_zip(job_id: str, job_result_dir: Path, file_results: list) -> Path:
    zip_path = OUTPUT_DIR / f"{job_id}.zip"
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zipped:
        for entry in file_results:
            if entry.get("status") == "success":
                output_path = Path(entry["output_path"])
                zipped.write(output_path, arcname=output_path.name)
        manifest = {
            "job_id": job_id,
            "generated_at": int(time.time()),
            "results": [
                {
                    "masked_name": entry["masked_name"],
                    "status": entry["status"],
                    "message": entry.get("message", ""),
                }
                for entry in file_results
            ],
        }
        zipped.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2))
    shutil.rmtree(job_result_dir, ignore_errors=True)
    return zip_path


def process_job(job_id: str) -> None:
    with _job_lock:
        job = _jobs[job_id]
        job["status"] = "running"

    result_dir = OUTPUT_DIR / job_id
    result_dir.mkdir(parents=True, exist_ok=True)
    file_results = []
    delete_rules = job.get("delete_rules", [])

    for index, file_id in enumerate(job["file_ids"], start=1):
        with _upload_lock:
            upload_info = _uploads[file_id]
        masked = upload_info["masked_name"]

        with _job_lock:
            job["files"][file_id]["status"] = "processing"
            job["files"][file_id]["progress"] = 60
            job["files"][file_id]["message"] = "处理中"
            job["processed"] = index - 1

        try:
            output_path = sanitize_document(
                Path(upload_info["assembled_path"]),
                upload_info["filename"],
                result_dir,
                delete_rules,
            )
            with _job_lock:
                job["files"][file_id]["status"] = "done"
                job["files"][file_id]["progress"] = 100
                job["files"][file_id]["message"] = "处理完成"
        except Exception as exc:  # noqa: BLE001
            with _job_lock:
                job["files"][file_id]["status"] = "failed"
                job["files"][file_id]["progress"] = 100
                job["files"][file_id]["message"] = str(exc)
            output_path = None

        file_results.append(
            {
                "file_id": file_id,
                "masked_name": masked,
                "status": "success" if output_path else "failed",
                "output_path": str(output_path) if output_path else "",
                "message": _jobs[job_id]["files"][file_id]["message"],
            }
        )

    with _job_lock:
        _jobs[job_id]["processed"] = len(job["file_ids"])

    try:
        zip_path = build_job_zip(job_id, result_dir, file_results)
    except Exception as exc:  # noqa: BLE001
        with _job_lock:
            _jobs[job_id]["results"] = file_results
            _jobs[job_id]["status"] = "failed"
            _jobs[job_id]["finished_at"] = int(time.time())
            _jobs[job_id]["error"] = str(exc)
        return

    with _job_lock:
        _jobs[job_id]["results"] = file_results
        _jobs[job_id]["zip_path"] = str(zip_path)
        _jobs[job_id]["status"] = "finished"
        _jobs[job_id]["finished_at"] = int(time.time())
        _jobs[job_id]["error"] = ""


@app.get("/")
def index():
    return render_template("index.html")


@app.post("/api/upload/init")
def upload_init():
    payload = request.get_json(silent=True) or {}
    filename = payload.get("filename", "").strip()
    size = int(payload.get("size", 0))
    chunk_size = int(payload.get("chunk_size", 0))
    total_chunks = int(payload.get("total_chunks", 0))

    if not filename or not allowed_file(filename):
        return jsonify({"error": "仅支持 .docx / .doc / .rtf 文件"}), 400
    if size <= 0 or size > MAX_FILE_SIZE:
        return jsonify({"error": "文件大小不合法"}), 400
    if chunk_size <= 0 or chunk_size > MAX_CHUNK_SIZE:
        return jsonify({"error": "分片大小不合法"}), 400
    if total_chunks <= 0:
        return jsonify({"error": "分片总数不合法"}), 400

    upload_id = uuid.uuid4().hex
    upload_path = UPLOAD_DIR / upload_id
    chunks_path = upload_path / "chunks"
    chunks_path.mkdir(parents=True, exist_ok=True)

    with _upload_lock:
        _uploads[upload_id] = {
            "id": upload_id,
            "filename": secure_filename(filename) or f"upload-{upload_id}",
            "masked_name": mask_filename(filename),
            "size": size,
            "chunk_size": chunk_size,
            "total_chunks": total_chunks,
            "received_chunks": set(),
            "received_bytes": 0,
            "state": "initiated",
            "upload_path": str(upload_path),
            "assembled_path": "",
            "created_at": int(time.time()),
        }

    return jsonify(
        {
            "upload_id": upload_id,
            "masked_name": _uploads[upload_id]["masked_name"],
        }
    )


@app.post("/api/upload/chunk")
def upload_chunk():
    upload_id = request.form.get("upload_id", "").strip()
    chunk_index_raw = request.form.get("chunk_index", "").strip()
    chunk_file = request.files.get("chunk")

    if not upload_id or upload_id not in _uploads:
        return jsonify({"error": "upload_id 无效"}), 400
    if chunk_file is None:
        return jsonify({"error": "缺少 chunk"}), 400
    if not chunk_index_raw.isdigit():
        return jsonify({"error": "chunk_index 非法"}), 400
    chunk_index = int(chunk_index_raw)

    with _upload_lock:
        upload_info = _uploads[upload_id]
    if chunk_index < 0 or chunk_index >= upload_info["total_chunks"]:
        return jsonify({"error": "chunk_index 超出范围"}), 400

    target_chunk = get_upload_chunk_path(upload_info, chunk_index)
    chunk_file.save(target_chunk)
    current_size = target_chunk.stat().st_size
    if current_size <= 0 or current_size > MAX_CHUNK_SIZE:
        target_chunk.unlink(missing_ok=True)
        return jsonify({"error": "chunk 大小不合法"}), 400

    with _upload_lock:
        received_chunks = upload_info["received_chunks"]
        if chunk_index not in received_chunks:
            received_chunks.add(chunk_index)
            upload_info["received_bytes"] += current_size
            upload_info["state"] = "uploading"

        progress = int((len(received_chunks) / upload_info["total_chunks"]) * 100)

    return jsonify(
        {
            "upload_id": upload_id,
            "progress": progress,
            "received_chunks": len(received_chunks),
            "total_chunks": upload_info["total_chunks"],
        }
    )


@app.post("/api/upload/complete")
def upload_complete():
    payload = request.get_json(silent=True) or {}
    upload_id = payload.get("upload_id", "").strip()
    if not upload_id or upload_id not in _uploads:
        return jsonify({"error": "upload_id 无效"}), 400

    with _upload_lock:
        upload_info = _uploads[upload_id]
        total_chunks = upload_info["total_chunks"]
        if len(upload_info["received_chunks"]) != total_chunks:
            return jsonify({"error": "文件分片未上传完整"}), 400

    assembled = Path(upload_info["upload_path"]) / upload_info["filename"]
    with assembled.open("wb") as merged:
        for index in range(total_chunks):
            chunk_path = get_upload_chunk_path(upload_info, index)
            if not chunk_path.exists():
                return jsonify({"error": f"分片缺失: {index}"}), 400
            with chunk_path.open("rb") as chunk_stream:
                shutil.copyfileobj(chunk_stream, merged, length=1024 * 1024)

    merged_size = assembled.stat().st_size
    if merged_size <= 0 or merged_size > MAX_FILE_SIZE or merged_size != upload_info["size"]:
        assembled.unlink(missing_ok=True)
        return jsonify({"error": "合并后的文件大小不合法"}), 400

    with _upload_lock:
        upload_info["assembled_path"] = str(assembled)
        upload_info["state"] = "completed"

    return jsonify(
        {
            "file_id": upload_id,
            "masked_name": upload_info["masked_name"],
            "size": merged_size,
        }
    )


@app.post("/api/jobs/start")
def start_job():
    payload = request.get_json(silent=True) or {}
    file_ids = payload.get("file_ids", [])
    raw_rules = payload.get("delete_rules", [])
    if not isinstance(file_ids, list) or not file_ids:
        return jsonify({"error": "file_ids 不能为空"}), 400
    if raw_rules is not None and not isinstance(raw_rules, list):
        return jsonify({"error": "delete_rules 必须为数组"}), 400

    try:
        delete_rules = normalize_delete_rules(raw_rules or [])
        compile_delete_rules(delete_rules)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except re.error as exc:
        return jsonify({"error": f"规则正则无效: {exc}"}), 400

    with _upload_lock:
        for file_id in file_ids:
            if file_id not in _uploads:
                return jsonify({"error": f"file_id 不存在: {file_id}"}), 400
            if _uploads[file_id]["state"] != "completed":
                return jsonify({"error": f"file_id 未完成上传: {file_id}"}), 400

    job_id = uuid.uuid4().hex
    files_meta = {}
    for file_id in file_ids:
        files_meta[file_id] = {
            "masked_name": _uploads[file_id]["masked_name"],
            "status": "queued",
            "progress": 0,
            "message": "等待处理",
        }

    with _job_lock:
        _jobs[job_id] = {
            "id": job_id,
            "file_ids": file_ids,
            "files": files_meta,
            "status": "queued",
            "processed": 0,
            "results": [],
            "zip_path": "",
            "delete_rules": delete_rules,
            "created_at": int(time.time()),
        }

    worker = threading.Thread(target=process_job, args=(job_id,), daemon=True)
    worker.start()
    return jsonify({"job_id": job_id})


@app.get("/api/jobs/<job_id>")
def get_job_status(job_id: str):
    with _job_lock:
        job = _jobs.get(job_id)
        if not job:
            return jsonify({"error": "job 不存在"}), 404

        files = []
        done = failed = 0
        for file_id, file_data in job["files"].items():
            state = file_data["status"]
            if state == "done":
                done += 1
            if state == "failed":
                failed += 1
            files.append(
                {
                    "file_id": file_id,
                    "masked_name": file_data["masked_name"],
                    "status": state,
                    "progress": file_data["progress"],
                    "message": file_data["message"],
                }
            )

        total = len(job["file_ids"])
        if total == 0:
            overall = 0
        else:
            overall = int(sum(item["progress"] for item in files) / total)

        return jsonify(
            {
                "job_id": job_id,
                "status": job["status"],
                "error": job.get("error", ""),
                "overall_progress": overall,
                "total": total,
                "done": done,
                "failed": failed,
                "files": files,
                "download_ready": bool(job["zip_path"]),
            }
        )


@app.get("/api/jobs/<job_id>/download")
def download_job_result(job_id: str):
    with _job_lock:
        job = _jobs.get(job_id)
        if not job:
            abort(404)
        zip_path = job.get("zip_path", "")
    if not zip_path:
        return jsonify({"error": "结果尚未生成"}), 400
    path = Path(zip_path)
    if not path.exists():
        return jsonify({"error": "结果文件不存在"}), 404
    return send_file(path, as_attachment=True, download_name=f"{job_id}.zip")


ensure_dirs()


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=8000, debug=False, threaded=True)
